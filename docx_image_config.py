import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image

# Directories
WORKSPACE = r"D:\ZoneMart"
MEDIA_DIR = os.path.join(WORKSPACE, "extracted_media")
USECASE_DIR = os.path.join(WORKSPACE, "extracted_usecase_images")
UI_DIR = os.path.join(WORKSPACE, "extracted_ui_images")
MD_PATH = r"C:\Users\kayy0\.gemini\antigravity\brain\ff1a3665-6bfa-430b-8f81-e5f1fb5e7295\BaoCao_TotNghiep_ZoneMart.md"
OUT_DOCX = os.path.join(WORKSPACE, "BaoCao_TotNghiep_ZoneMart.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def make_callout_box(doc, title, content_lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F6F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    border_color = "2E5B82"
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(title)
    run_title.bold = True
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(11.5)
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    for line in content_lines:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15
        run = p2.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def insert_image_with_caption(doc, img_path, caption_text):
    if not os.path.exists(img_path):
        print(f"[WARN] Image not found: {img_path}")
        return False
        
    try:
        # Check image aspect ratio
        with Image.open(img_path) as img:
            w, h = img.size
            aspect = h / w if w > 0 else 1.0
            
        # Target width
        max_w_in = 6.0
        # If very tall, limit height to 6.2 inches
        if max_w_in * aspect > 6.2:
            target_w = Inches(6.2 / aspect)
        else:
            target_w = Inches(max_w_in)
            
        # Add image paragraph
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.keep_with_next = True
        
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=target_w)
        
        # Add caption paragraph
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Highlight figure label (e.g. Hình 2.1:)
        m = re.match(r'^(Hình\s+[\d\.\w]+:)(.*)$', caption_text.strip())
        if m:
            r1 = p_cap.add_run(m.group(1))
            r1.bold = True
            r1.italic = True
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(11)
            r1.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            
            r2 = p_cap.add_run(m.group(2))
            r2.italic = True
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(11)
            r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        else:
            r = p_cap.add_run(caption_text.strip())
            r.italic = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            
        return True
    except Exception as e:
        print(f"[ERR] Failed inserting image {img_path}: {e}")
        return False

# Mapping from keywords in markdown lines to image paths
IMAGE_MAPPING = {
    # Chapter 1
    "Hình 1.1": os.path.join(MEDIA_DIR, "1.1_mo_hinh_tong_the_zonemart.jpg"),
    "Hình 1.2": os.path.join(MEDIA_DIR, "1.2_mo_hinh_quan_ly_chi_nhanh_zonemart.jpg"),
    "Hình 1.3": os.path.join(MEDIA_DIR, "1.3_chuoi_gia_tri_nghiep_vu_zonemart.jpg"),
    
    # Chapter 2 - Business Flows
    "Hình 2.1": os.path.join(MEDIA_DIR, "2.1_tong_quan_luong_nghiep_vu_zonemart.jpg"),
    "Hình 2.2a": os.path.join(MEDIA_DIR, "2.2_luong_dang_ky_gian_hang_seller.jpg"),
    "Hình 2.2b": os.path.join(MEDIA_DIR, "2.3_luong_mua_hang_ai_tu_van.jpg"),
    "Hình 2.2c": os.path.join(MEDIA_DIR, "2.4_luong_xu_ly_don_hang_store.jpg"),
    "Hình 2.2d": os.path.join(MEDIA_DIR, "2.5_luong_giao_hang_hoa_toc_gps.jpg"),
    "Hình 2.2e": os.path.join(MEDIA_DIR, "2.6_luong_truy_xuat_attp_khiem_soat.jpg"),
    "Hình 2.2f": os.path.join(MEDIA_DIR, "3.1_tong_hop_mo_hinh_phan_tich_zonemart.jpg"),
    "Hình 2.2g": os.path.join(MEDIA_DIR, "3.2_phan_tich_yeu_cau_chi_tiet_zonemart.jpg"),
    
    # Chapter 2 - Use Cases
    "Hình 2.3:": os.path.join(MEDIA_DIR, "4.1_use_case_tong_quat_zonemart.jpg"),
    "Hình 2.3a": os.path.join(MEDIA_DIR, "4.2_use_case_khach_hang_zonemart.jpg"),
    "Hình 2.3b": os.path.join(MEDIA_DIR, "4.3_use_case_chu_gian_hang_zonemart.jpg"),
    "Hình 2.3c": os.path.join(MEDIA_DIR, "4.4_use_case_tai_xe_shipper_zonemart.jpg"),
    "Hình 2.3d": os.path.join(MEDIA_DIR, "4.5_use_case_quan_tri_admin_zonemart.jpg"),
    
    "Hình 2.4:": os.path.join(USECASE_DIR, "hinh_2_4.png"),
    "Hình 2.5:": os.path.join(USECASE_DIR, "hinh_2_5.png"),
    "Hình 2.6:": os.path.join(USECASE_DIR, "hinh_2_6.png"),
    "Hình 2.7:": os.path.join(USECASE_DIR, "hinh_2_7.png"),
    "Hình 2.8:": os.path.join(USECASE_DIR, "hinh_2_8.png"),
    "Hình 2.9:": os.path.join(USECASE_DIR, "hinh_2_9.png"),
    "Hình 2.10:": os.path.join(USECASE_DIR, "hinh_2_10.png"),
    "Hình 2.11:": os.path.join(USECASE_DIR, "hinh_2_11.png"),
    "Hình 2.12:": os.path.join(USECASE_DIR, "hinh_2_12.png"),
    "Hình 2.13:": os.path.join(USECASE_DIR, "hinh_2_13.png"),
    "Hình 2.14:": os.path.join(USECASE_DIR, "hinh_2_14.png"),
    "Hình 2.15:": os.path.join(USECASE_DIR, "hinh_2_15.png"),
    "Hình 2.16:": os.path.join(USECASE_DIR, "hinh_2_16.png"),
    "Hình 2.17:": os.path.join(USECASE_DIR, "hinh_2_17.png"),
    "Hình 2.18:": os.path.join(USECASE_DIR, "hinh_2_18.png"),
    "Hình 2.19:": os.path.join(USECASE_DIR, "hinh_2_19.png"),
    "Hình 2.20:": os.path.join(USECASE_DIR, "hinh_2_20.png"),
    
    # Chapter 2 - Activity Diagrams
    "Hình 2.21a": os.path.join(MEDIA_DIR, "4.6_activity_dang_ky_duyet_store.jpg"),
    "Hình 2.21b": os.path.join(MEDIA_DIR, "4.7_activity_mua_hang_ai_checkout.jpg"),
    "Hình 2.21c": os.path.join(MEDIA_DIR, "4.8_activity_dieu_phoi_shipper_gps.jpg"),
    "Hình 2.21d": os.path.join(MEDIA_DIR, "4.9_activity_khieu_nai_truy_xuat_attp.jpg"),
    
    # Chapter 2 - Database ERD
    "Hình 2.22:": os.path.join(MEDIA_DIR, "5.1_erd_tong_the_csdl_zonemart.jpg"),
    "Hình 2.22a": os.path.join(MEDIA_DIR, "5.2_erd_phan_he_nguoi_dung_faceid.jpg"),
    "Hình 2.22b": os.path.join(MEDIA_DIR, "5.3_erd_phan_he_gian_hang_store.jpg"),
    "Hình 2.22c": os.path.join(MEDIA_DIR, "5.4_erd_phan_he_san_pham_ai_attp.jpg"),
    "Hình 2.22d": os.path.join(MEDIA_DIR, "5.5_erd_phan_he_gio_hang_don_hang.jpg"),
    "Hình 2.22e": os.path.join(MEDIA_DIR, "5.6_erd_phan_he_logistics_shipper_gps.jpg"),
    
    # Chapter 3 - System Architecture & UI Overview
    "Hình 3.1a": os.path.join(MEDIA_DIR, "6.1_kien_truc_tich_hop_he_thong_ngoai.jpg"),
    "Hình 3.1b": os.path.join(MEDIA_DIR, "6.2_kien_truc_ung_dung_modular_monolith.jpg"),
    "Hình 3.2a": os.path.join(MEDIA_DIR, "7.1_thiet_ke_giao_dien_tong_quan_ui.jpg"),
    "Hình 3.2b": os.path.join(MEDIA_DIR, "7.2_thiet_ke_giao_dien_nang_cao.jpg"),
    
    # Chapter 3 - UI Screenshots
    "Hình 3.1:": os.path.join(UI_DIR, "hinh_3_1.png"),
    "Hình 3.2:": os.path.join(UI_DIR, "hinh_3_2.png"),
    "Hình 3.3:": os.path.join(UI_DIR, "hinh_3_3.png"),
    "Hình 3.4:": os.path.join(UI_DIR, "hinh_3_4.png"),
    "Hình 3.5:": os.path.join(UI_DIR, "hinh_3_5.png"),
    "Hình 3.6:": os.path.join(UI_DIR, "hinh_3_6.png"),
    "Hình 3.7:": os.path.join(UI_DIR, "hinh_3_7.png"),
    "Hình 3.8:": os.path.join(UI_DIR, "hinh_3_8.png"),
    "Hình 3.9:": os.path.join(UI_DIR, "hinh_3_9.png"),
    "Hình 3.10:": os.path.join(UI_DIR, "hinh_3_10.png"),
    "Hình 3.11:": os.path.join(UI_DIR, "hinh_3_11.png"),
    "Hình 3.12:": os.path.join(UI_DIR, "hinh_3_12.png"),
    "Hình 3.13:": os.path.join(UI_DIR, "hinh_3_13.png"),
    "Hình 3.14:": os.path.join(UI_DIR, "hinh_3_14.png"),
    "Hình 3.15:": os.path.join(UI_DIR, "hinh_3_15.png"),
    "Hình 3.16:": os.path.join(UI_DIR, "hinh_3_16.png"),
    "Hình 3.17:": os.path.join(UI_DIR, "hinh_3_17.png"),
    "Hình 3.18:": os.path.join(UI_DIR, "hinh_3_18.png")
}

def find_matched_image(caption_line):
    clean_line = caption_line.strip().replace('**', '').replace('*', '')
    for key, path in IMAGE_MAPPING.items():
        if key in clean_line:
            if os.path.exists(path):
                return path, clean_line
    return None, clean_line

print(f"[INIT] Loaded {len(IMAGE_MAPPING)} image mappings.")

